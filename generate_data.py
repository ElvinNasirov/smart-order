"""
generate_data.py — 3 синтетических Excel-отчёта одного маркета в Баку, каждый в своём формате.

  sales_2021_06.xlsx  июнь 2021, колонки на азербайджанском, дата ДД.ММ.ГГГГ, строка = товар за день
  sales_2022_03.xlsx  март 2022, колонки на русском, дата "5 мар 2022", 3 строки шапки, цена с "AZN"
  sales_2024_10.xlsx  октябрь 2024, широкий формат (товары x дни 1..31), цены на отдельном листе

  ground_truth.csv     чистые данные всех 3 файлов: date, product_name, category, quantity, unit_price
                       (product_name и category — канонические, на английском)
  errors_injected.csv  список внесённых ошибок: file, sheet, excel_row, column, type, details

Запуск: python generate_data.py
"""
import calendar
import copy
import csv
import random
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

random.seed(42)

# ----------------------------------------------------------------------------
# Каталог товаров: английское (каноническое) имя, az, ru, категории, единица, цена AZN, базовый спрос/день
# ----------------------------------------------------------------------------
_RAW = [
    ("Milk 1L",          "Süd 1L",            "Молоко 1л",               "Dairy",     "Süd məhsulları", "pcs", 1.60, 45, []),
    ("Kefir",            "Kefir",             "Кефир",                   "Dairy",     "Süd məhsulları", "pcs", 1.40, 25, []),
    ("Cheese",           "Pendir",            "Сыр",                     "Dairy",     "Süd məhsulları", "kg",  12.0,  8, []),
    ("Butter",           "Kərə yağı",         "Сливочное масло",         "Dairy",     "Süd məhsulları", "pcs", 4.50, 12, []),
    ("Eggs (10 pcs)",    "Yumurta (10 əd.)",  "Яйца (10 шт.)",           "Dairy",     "Süd məhsulları", "pcs", 3.20, 30, []),
    ("Bread",            "Çörək",             "Хлеб",                    "Bakery",    "Çörək məmulatları", "pcs", 0.80, 140, []),
    ("Lavash",           "Lavaş",             "Лаваш",                   "Bakery",    "Çörək məmulatları", "pcs", 0.70, 60, []),
    ("Tandoor bread",    "Təndir çörəyi",     "Тандырный хлеб",          "Bakery",    "Çörək məmulatları", "pcs", 0.90, 70, []),
    ("Black tea",        "Qara çay",          "Чёрный чай",              "Beverages", "İçkilər",        "pcs", 4.80, 18, ["tea"]),
    ("Ice cream",        "Dondurma",          "Мороженое",               "Frozen",    "Dondurulmuş",    "pcs", 1.20, 40, ["cold"]),
    ("Mineral water",    "Mineral su",        "Минеральная вода",        "Beverages", "İçkilər",        "pcs", 0.70, 80, ["cold_mild"]),
    ("Cola",             "Kola",              "Кола",                    "Beverages", "İçkilər",        "pcs", 1.50, 50, ["cold_mild"]),
    ("Juice",            "Meyvə şirəsi",      "Сок",                     "Beverages", "İçkilər",        "pcs", 2.20, 35, ["cold_mild"]),
    ("Beef",             "Mal əti",           "Говядина",                "Meat",      "Ət məhsulları",  "kg",  14.0, 18, ["meat"]),
    ("Chicken",          "Toyuq əti",         "Курица",                  "Meat",      "Ət məhsulları",  "kg",  6.50, 30, ["meat"]),
    ("Lamb",             "Quzu əti",          "Баранина",                "Meat",      "Ət məhsulları",  "kg",  16.0,  8, ["meat"]),
    ("Sausage",          "Kolbasa",           "Колбаса",                 "Meat",      "Ət məhsulları",  "kg",  9.00, 12, ["meat"]),
    ("Chocolate",        "Şokolad",           "Шоколад",                 "Sweets",    "Şirniyyat",      "pcs", 2.50, 30, ["sweet"]),
    ("Candy",            "Konfet",            "Конфеты",                 "Sweets",    "Şirniyyat",      "kg",  8.00, 12, ["sweet"]),
    ("Baklava",          "Paxlava",           "Пахлава",                 "Sweets",    "Şirniyyat",      "kg",  18.0,  6, ["sweet"]),
    ("Shakarbura",       "Şəkərbura",         "Шекербура",               "Sweets",    "Şirniyyat",      "kg",  15.0,  3, ["sweet"]),
    ("Rice",             "Düyü",              "Рис",                     "Grocery",   "Ərzaq",          "kg",  3.20, 30, []),
    ("Sugar",            "Şəkər",             "Сахар",                   "Grocery",   "Ərzaq",          "kg",  1.60, 40, []),
    ("Sunflower oil 1L", "Günəbaxan yağı 1L", "Подсолнечное масло 1л",  "Grocery",   "Ərzaq",          "pcs", 4.20, 22, []),
    ("Flour",            "Un",                "Мука",                    "Grocery",   "Ərzaq",          "kg",  1.10, 35, []),
]
KEYS = ["en", "az", "ru", "cat_en", "cat_az", "unit", "price", "base", "tags"]
PRODUCTS = [dict(zip(KEYS, r)) for r in _RAW]

# сезонность (множитель по месяцам)
COLD = {1: .25, 2: .25, 3: .4, 4: .6, 5: .9, 6: 1.6, 7: 2.0, 8: 2.1, 9: 1.3, 10: .7, 11: .4, 12: .3}
COLD_MILD = {m: 1 + (v - 1) * 0.4 for m, v in COLD.items()}
TEA = {1: 1.4, 2: 1.35, 3: 1.2, 4: 1.0, 5: .9, 6: .8, 7: .75, 8: .75, 9: .9, 10: 1.1, 11: 1.25, 12: 1.4}
# Новруз (14–24 марта, пик 19-го): прирост спроса
NOVRUZ = {"Shakarbura": 3.0, "Baklava": 2.0, "Flour": 1.0, "Eggs (10 pcs)": 0.8, "Candy": 0.8,
          "Sugar": 0.6, "Rice": 0.5, "Chocolate": 0.5}


def mult(p, d):
    """Множитель спроса для товара p в день d: сезон, праздники, выходные."""
    m = 1.0
    tags = p["tags"]
    if "cold" in tags:
        m *= COLD[d.month]
    if "cold_mild" in tags:
        m *= COLD_MILD[d.month]
    if "tea" in tags:
        m *= TEA[d.month]
    # Новый год: рост сладостей и мяса
    ny = (d.month == 12 and d.day >= 24) or (d.month == 1 and d.day <= 2)
    pre_ny = d.month == 12 and 15 <= d.day < 24
    if "sweet" in tags or "meat" in tags:
        if ny:
            m *= 2.2
        elif pre_ny:
            m *= 1.4
    # Novruz
    if d.month == 3 and 14 <= d.day <= 24 and p["en"] in NOVRUZ:
        m *= 1 + NOVRUZ[p["en"]] * max(0, 1 - abs(d.day - 19) / 6)
    # День восстановления независимости (17–18 окт): небольшой рост
    if d.month == 10 and d.day in (17, 18) and ("sweet" in tags or "meat" in tags):
        m *= 1.3
    # выходные
    if d.weekday() >= 5:
        m *= 1.2
    return m


def month_data(year, month):
    """qty[en][day], price[en] за месяц."""
    last = calendar.monthrange(year, month)[1]
    infl = 1 + 0.06 * (year - 2021)
    qty, price = {}, {}
    for p in PRODUCTS:
        price[p["en"]] = round(round(p["price"] * infl * random.uniform(0.97, 1.03) * 20) / 20, 2)
        qty[p["en"]] = {}
        for day in range(1, last + 1):
            d = date(year, month, day)
            x = p["base"] * mult(p, d) * random.lognormvariate(0, 0.15)
            qty[p["en"]][day] = max(0.5, round(x, 1)) if p["unit"] == "kg" else max(1, int(round(x)))
    return qty, price


def typo(s):
    """Опечатка: пропускаем одну букву в середине."""
    i = random.randint(1, len(s) - 2)
    return s[:i] + s[i + 1:]


# ----------------------------------------------------------------------------
# стили Excel
# ----------------------------------------------------------------------------
HDR_FONT = Font(bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor="305496")


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HDR_FONT
        cell.fill = HDR_FILL
        cell.alignment = Alignment(horizontal="center")


def set_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ----------------------------------------------------------------------------
# длинные форматы (файлы 1 и 2): общая сборка и внесение ошибок
# ----------------------------------------------------------------------------
def build_long(year, month, lang, date_fmt):
    qty, price = month_data(year, month)
    rows = []
    for day in range(1, calendar.monthrange(year, month)[1] + 1):
        d = date(year, month, day)
        for p in PRODUCTS:
            rows.append({
                "d": d, "p": p, "q": qty[p["en"]][day], "pr": price[p["en"]],
                "date_out": date_fmt(d), "name_out": p[lang],
                "cat_out": p["cat_az"] if lang == "az" else p["cat_en"],
                "qty_out": qty[p["en"]][day], "price_out": price[p["en"]],
                "truth": True, "errs": [],
            })
    return rows


def inject_long(rows, cols, milk_typo, text_value, invalid_date, has_cat):
    """cols: заголовки колонок для записи в журнал ошибок."""
    used = set()

    def pick(cond=lambda r: True):
        r = random.choice([x for x in rows if id(x) not in used and cond(x)])
        used.add(id(r))
        return r

    # 1) опечатка/другой язык в названии молока
    r = pick(lambda r: r["p"]["en"] == "Milk 1L")
    old = r["name_out"]
    r["name_out"] = milk_typo
    r["errs"].append((cols["name"], "typo_name", f"{old} -> {milk_typo}"))
    # 2) обычная опечатка
    r = pick(lambda r: len(r["name_out"]) >= 6 and r["p"]["en"] != "Milk 1L")
    old = r["name_out"]
    r["name_out"] = typo(old)
    r["errs"].append((cols["name"], "typo_name", f"{old} -> {r['name_out']}"))
    # 3) пустое количество
    r = pick()
    r["errs"].append((cols["qty"], "empty_cell", f"true value {r['q']}"))
    r["qty_out"] = None
    # 4) пустая цена
    r = pick()
    r["errs"].append((cols["price"], "empty_cell", f"true value {r['pr']}"))
    r["price_out"] = None
    # 5) текст вместо числа
    r = pick()
    r["errs"].append((cols["qty"], "text_instead_of_number", f"true value {r['q']} -> '{text_value}'"))
    r["qty_out"] = text_value
    # 6) отрицательное количество
    r = pick()
    r["errs"].append((cols["qty"], "negative_quantity", f"true value {r['q']} -> {-r['q']}"))
    r["qty_out"] = -r["q"]
    # 7) невозможная дата
    r = pick(lambda r: r["d"].day == 30)
    new = invalid_date(r["d"])
    r["errs"].append((cols["date"], "invalid_date", f"true {r['d'].isoformat()} -> '{new}'"))
    r["date_out"] = new
    # 8) пустая категория
    if has_cat:
        r = pick()
        r["errs"].append((cols["cat"], "empty_cell", f"true value {r['cat_out']}"))
        r["cat_out"] = None
    # 9) дубль строки (вставляется сразу после оригинала)
    r = pick()
    idx = next(i for i, x in enumerate(rows) if x is r)
    dup = copy.deepcopy(r)
    dup["truth"] = False
    dup["errs"] = [("(вся строка)", "duplicate_row", "точная копия предыдущей строки")]
    rows.insert(idx + 1, dup)
    return rows


def truth_from_long(rows):
    return [(r["d"].isoformat(), r["p"]["en"], r["p"]["cat_en"], r["q"], r["pr"]) for r in rows if r["truth"]]


def errors_from_long(rows, fname, sheet, first_data_row):
    out = []
    for i, r in enumerate(rows):
        for col, typ, det in r["errs"]:
            out.append((fname, sheet, first_data_row + i, col, typ, det))
    return out


# ----------------------------------------------------------------------------
# Файл 1: sales_2021_06.xlsx (азербайджанский)
# ----------------------------------------------------------------------------
def make_file1():
    fname = "sales_2021_06.xlsx"
    rows = build_long(2021, 6, "az", lambda d: d.strftime("%d.%m.%Y"))
    cols = {"date": "Tarix", "name": "Məhsul", "cat": "Kateqoriya", "qty": "Miqdar", "price": "Qiymət"}
    rows = inject_long(rows, cols, milk_typo="Молоко", text_value="beş",
                       invalid_date=lambda d: f"31.{d.month:02d}.{d.year}", has_cat=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Satış"
    ws.append(["Tarix", "Məhsul", "Kateqoriya", "Miqdar", "Qiymət"])
    style_header(ws, 1, 5)
    for r in rows:
        ws.append([r["date_out"], r["name_out"], r["cat_out"], r["qty_out"], r["price_out"]])
    set_widths(ws, [13, 22, 20, 10, 10])
    ws.freeze_panes = "A2"
    wb.save(fname)
    return truth_from_long(rows), errors_from_long(rows, fname, "Satış", 2), len(rows)


# ----------------------------------------------------------------------------
# Файл 2: sales_2022_03.xlsx (русский, 3 строки шапки, "5 мар 2022", цена с AZN)
# ----------------------------------------------------------------------------
RU_MON = {1: "янв", 2: "фев", 3: "мар", 4: "апр", 5: "мая", 6: "июн",
          7: "июл", 8: "авг", 9: "сен", 10: "окт", 11: "нояб", 12: "дек"}


def make_file2():
    fname = "sales_2022_03.xlsx"
    rows = build_long(2022, 3, "ru", lambda d: f"{d.day} {RU_MON[d.month]} {d.year}")
    cols = {"date": "Дата", "name": "Наименование", "cat": None, "qty": "Кол-во", "price": "Цена за ед."}
    rows = inject_long(rows, cols, milk_typo="Moloko", text_value="пять",
                       invalid_date=lambda d: f"30 фев {d.year}", has_cat=False)
    wb = Workbook()
    ws = wb.active
    ws.title = "Продажи март"
    ws.append(["Магазин «Баку Маркет», филиал Насими, г. Баку, ул. Рашид Бейбутов 12"])
    ws.append(["Отчёт о продажах за март 2022 г."])
    ws.append(["Ответственный: Алиев Р.   |   Валюта: AZN"])
    for r_ in (1, 2, 3):
        ws.merge_cells(start_row=r_, start_column=1, end_row=r_, end_column=5)
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"].font = Font(bold=True)
    ws.append(["Дата", "Наименование", "Кол-во", "Цена за ед.", "Сумма"])
    style_header(ws, 4, 5)
    for r in rows:
        price_txt = None if r["price_out"] is None else f"{r['price_out']:.2f} AZN"
        total = round(r["q"] * r["pr"], 2)  # сумма считается по истинным значениям
        ws.append([r["date_out"], r["name_out"], r["qty_out"], price_txt, total])
    set_widths(ws, [14, 26, 10, 14, 12])
    wb.save(fname)
    return truth_from_long(rows), errors_from_long(rows, fname, "Продажи март", 5), len(rows)


# ----------------------------------------------------------------------------
# Файл 3: sales_2024_10.xlsx (широкий формат, цены на отдельном листе)
# ----------------------------------------------------------------------------
def make_file3():
    fname = "sales_2024_10.xlsx"
    year, month = 2024, 10
    last = calendar.monthrange(year, month)[1]
    qty, price = month_data(year, month)
    rows = [{"p": p, "name": p["en"], "true": qty[p["en"]], "cells": dict(qty[p["en"]]),
             "truth": True, "errs": []} for p in PRODUCTS]

    milk = next(r for r in rows if r["p"]["en"] == "Milk 1L")
    others = [r for r in rows if r is not milk]
    sel = random.sample(others, 7)

    old = milk["name"]
    milk["name"] = "Moloko 1L"
    milk["errs"].append(("A", "typo_name", f"{old} -> {milk['name']}"))

    r = sel[0]
    old = r["name"]
    r["name"] = typo(old)
    r["errs"].append(("A", "typo_name", f"{old} -> {r['name']}"))

    def cell_error(r, kind, new, det_fmt):
        day = random.randint(1, last)
        col = get_column_letter(2 + day)  # A=Product, B=Category, C=день 1
        r["errs"].append((col, kind, det_fmt.format(day=day, true=r["true"][day], new=new)))
        r["cells"][day] = new

    cell_error(sel[2], "empty_cell", None, "day {day}: true value {true}")
    cell_error(sel[3], "empty_cell", None, "day {day}: true value {true}")
    cell_error(sel[4], "text_instead_of_number", "n/a", "day {day}: true value {true} -> '{new}'")
    cell_error(sel[5], "text_instead_of_number", "five", "day {day}: true value {true} -> '{new}'")
    # отрицательное
    r = sel[6]
    day = random.randint(1, last)
    r["errs"].append((get_column_letter(2 + day), "negative_quantity",
                      f"day {day}: true value {r['true'][day]} -> {-r['true'][day]}"))
    r["cells"][day] = -r["true"][day]

    # дубль всей строки товара
    r = sel[1]
    idx = next(i for i, x in enumerate(rows) if x is r)
    dup = copy.deepcopy(r)
    dup["truth"] = False
    dup["errs"] = [("(вся строка)", "duplicate_row", "точная копия предыдущей строки товара")]
    rows.insert(idx + 1, dup)

    wb = Workbook()
    ws = wb.active
    ws.title = "Sales Oct 2024"
    ws.append(["Daily sales report — October 2024 (quantity sold per day)"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    ws.append(["Product", "Category"] + list(range(1, last + 1)))
    style_header(ws, 3, 2 + last)
    for r in rows:
        ws.append([r["name"], r["p"]["cat_en"]] + [r["cells"][d] for d in range(1, last + 1)])
    set_widths(ws, [22, 12] + [5] * last)
    ws.freeze_panes = "C4"

    wp = wb.create_sheet("Prices")
    wp.append(["Product", "Unit price (AZN)"])
    style_header(wp, 1, 2)
    for p in PRODUCTS:
        wp.append([p["en"], price[p["en"]]])
    set_widths(wp, [22, 16])
    wb.save(fname)

    truth = []
    for r in rows:
        if r["truth"]:
            for day in range(1, last + 1):
                truth.append((date(year, month, day).isoformat(), r["p"]["en"], r["p"]["cat_en"],
                              r["true"][day], price[r["p"]["en"]]))
    errs = []
    for i, r in enumerate(rows):
        for col, typ, det in r["errs"]:
            errs.append((fname, "Sales Oct 2024", 4 + i, col, typ, det))
    return truth, errs, len(rows)


# ----------------------------------------------------------------------------
if __name__ == "__main__":
    truth, errs = [], []
    for fn in (make_file1, make_file2, make_file3):
        t, e, n = fn()
        truth += t
        errs += e
        print(f"{fn.__name__}: {n} строк в Excel, {len(t)} чистых записей, {len(e)} ошибок")

    truth.sort(key=lambda x: (x[0], x[1]))
    with open("ground_truth.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["date", "product_name", "category", "quantity", "unit_price"])
        w.writerows(truth)
    with open("errors_injected.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["file", "sheet", "excel_row", "column", "type", "details"])
        w.writerows(errs)
    print(f"ground_truth.csv: {len(truth)} строк; errors_injected.csv: {len(errs)} строк")
