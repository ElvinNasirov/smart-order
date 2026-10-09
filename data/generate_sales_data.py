#!/usr/bin/env python3
"""Synthetic sales reports for ONE Baku market, kept by three different "people".

  sales_2022_12.xlsx  Azerbaijani headers, real dates DD.MM.YYYY, long format
  sales_2022_11.xlsx  Russian headers, 3 title rows, ISO text dates, "1.95 AZN" prices
  sales_2022_02.xlsx  wide format (products x days 1..31), prices on a second sheet
  ground_truth.csv    clean data of all 3 files: date, product_name, category, quantity, unit_price
  errors_injected.csv every injected error: file, sheet, excel_row, column, error_type, ...

Usage: python generate_sales_data.py [out_dir]
"""
import calendar
import copy
import datetime as dt
import random
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUT = Path(sys.argv[1] if len(sys.argv) > 1 else "output")
OUT.mkdir(parents=True, exist_ok=True)
SEED = 42
rng = np.random.default_rng(SEED)
rnd = random.Random(SEED)

# ----------------------------------------------------------------- catalogue
# category key -> (Russian canonical, Azerbaijani)
CATS = {
    "dairy": ("Молочные продукты", "Süd məhsulları"),
    "bakery": ("Хлеб и выпечка", "Çörək məmulatları"),
    "frozen": ("Замороженные продукты", "Dondurulmuş məhsullar"),
    "drinks": ("Напитки", "İçkilər"),
    "meat": ("Мясо", "Ət məhsulları"),
    "sweets": ("Сладости", "Şirniyyat"),
    "grocery": ("Бакалея", "Ərzaq"),
}
# ru (canonical), az, az without diacritics (Feb file), category, base price AZN, base daily qty
PRODUCTS = [
    dict(ru=a, az=b, az2=c, cat=d, price=e, base=f)
    for a, b, c, d, e, f in [
        ("Молоко", "Süd", "Sud", "dairy", 1.80, 38),
        ("Кефир", "Kefir", "Kefir", "dairy", 1.60, 14),
        ("Сыр", "Pendir", "Pendir", "dairy", 12.50, 8),
        ("Йогурт", "Yoqurt", "Yoqurt", "dairy", 1.20, 20),
        ("Яйца", "Yumurta", "Yumurta", "dairy", 3.20, 26),
        ("Хлеб", "Çörək", "Corek", "bakery", 0.80, 110),
        ("Лаваш", "Lavaş", "Lavas", "bakery", 0.60, 55),
        ("Мороженое", "Dondurma", "Dondurma", "frozen", 1.50, 22),
        ("Пельмени", "Pelmeni", "Pelmeni", "frozen", 6.00, 8),
        ("Чай", "Çay", "Cay", "drinks", 3.50, 16),
        ("Кофе", "Qəhvə", "Qehve", "drinks", 8.00, 9),
        ("Сок", "Şirə", "Sire", "drinks", 2.80, 20),
        ("Вода", "Su", "Su", "drinks", 0.60, 60),
        ("Говядина", "Mal əti", "Mal eti", "meat", 14.00, 11),
        ("Курица", "Toyuq əti", "Toyuq eti", "meat", 7.50, 20),
        ("Колбаса", "Kolbasa", "Kolbasa", "meat", 9.00, 12),
        ("Баранина", "Quzu əti", "Quzu eti", "meat", 16.00, 6),
        ("Шоколад", "Şokolad", "Sokolad", "sweets", 2.50, 30),
        ("Конфеты", "Konfet", "Konfet", "sweets", 6.00, 15),
        ("Печенье", "Peçenye", "Pecenye", "sweets", 2.20, 25),
        ("Халва", "Halva", "Halva", "sweets", 4.50, 8),
        ("Рис", "Düyü", "Duyu", "grocery", 2.80, 15),
        ("Сахар", "Şəkər", "Seker", "grocery", 1.50, 25),
        ("Макароны", "Makaron", "Makaron", "grocery", 1.30, 20),
        ("Масло подсолнечное", "Günəbaxan yağı", "Gunebaxan yagi", "grocery", 3.80, 12),
    ]
]
N = len(PRODUCTS)

# ----------------------------------------------------------------- seasonality
ICE = {1: .15, 2: .20, 3: .35, 4: .60, 5: 1.0, 6: 1.8, 7: 2.6, 8: 3.0, 9: 1.8, 10: .90, 11: .35, 12: .30}
COLD = {1: .8, 2: .8, 3: .9, 4: 1.0, 5: 1.2, 6: 1.4, 7: 1.6, 8: 1.6, 9: 1.3, 10: 1.0, 11: .85, 12: .8}
TEA = {1: 1.35, 2: 1.3, 3: 1.15, 4: 1.0, 5: .85, 6: .75, 7: .7, 8: .7, 9: .85, 10: 1.0, 11: 1.2, 12: 1.35}
PRICE_FACTOR = {2: 1.00, 11: 1.09, 12: 1.10}  # mild inflation between months


def season(p, d):
    m, k = d.month, 1.0
    if p["ru"] == "Мороженое":
        k = ICE[m]
    elif p["ru"] in ("Сок", "Вода"):
        k = COLD[m]
    elif p["ru"] == "Чай":
        k = TEA[m]
    if m == 12 and p["cat"] in ("sweets", "meat"):  # New Year ramp, peak on 29-31 Dec
        k *= 1 + 1.2 * ((d.day - 1) / 30) ** 1.5 + (0.4 if d.day >= 29 else 0)
    if d.weekday() >= 5:
        k *= 1.2
    return k


def generate(year, month):
    days = calendar.monthrange(year, month)[1]
    f = PRICE_FACTOR[month]
    prices = [round(round(p["price"] * f * 20) / 20, 2) for p in PRODUCTS]
    qty = np.array([[rng.poisson(p["base"] * season(p, dt.date(year, month, d))) for p in PRODUCTS]
                    for d in range(1, days + 1)])
    return days, prices, qty


def truth_rows(year, month, days, prices, qty):
    return [dict(date=dt.date(year, month, d + 1).isoformat(), product_name=p["ru"],
                 category=CATS[p["cat"]][0], quantity=int(qty[d][i]), unit_price=prices[i])
            for d in range(days) for i, p in enumerate(PRODUCTS) if qty[d][i] > 0]


# ----------------------------------------------------------------- typos
_RU = dict(zip("абвгдеёжзийклмнопрстуфхцчшщъыьэюя",
               ["a", "b", "v", "g", "d", "e", "e", "zh", "z", "i", "y", "k", "l", "m", "n", "o", "p", "r", "s",
                "t", "u", "f", "kh", "ts", "ch", "sh", "sch", "", "y", "", "e", "yu", "ya"]))
_AZ = str.maketrans("əƏıİöÖüÜçÇşŞğĞ", "eEiIoOuUcCsSgG")


def translit(s):
    out = "".join(_RU.get(c.lower(), c) for c in s)
    return out[:1].upper() + out[1:]


def make_typo(name, lang):
    opts = {translit(name) if lang == "ru" else name.translate(_AZ)}
    if len(name) >= 4:
        i = rnd.randrange(1, len(name) - 1)
        opts.add(name[:i] + name[i + 1:])        # dropped letter
        opts.add(name[:i] + name[i] + name[i:])  # doubled letter
    opts.update({name.upper(), name + " "})       # CAPS / trailing space
    opts.discard(name)
    return rnd.choice(sorted(opts))


# ----------------------------------------------------------------- helpers
FONT, BOLD = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
ERRORS, ROWMAP = [], {}


def put(ws, r, c, v, bold=False, fmt=None, align=None):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font = BOLD if bold else FONT
    if fmt:
        cell.number_format = fmt
    if align:
        cell.alignment = Alignment(horizontal=align)
    return cell


def header(ws, r, names):
    for c, n in enumerate(names, 1):
        put(ws, r, c, n, True, align="center").fill = PatternFill("solid", fgColor="D9E1F2")


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def s(x):
    if x is None:
        return "(empty)"
    if isinstance(x, dt.date):
        return x.strftime("%d.%m.%Y")
    if isinstance(x, float):
        return f"{x:.2f}"
    return str(x)


def log(file, sheet, rid, column, etype, orig, inj, date, product):
    ERRORS.append(dict(file=file, sheet=sheet, _id=rid, column=column, error_type=etype,
                       original_value=s(orig), injected_value=s(inj), true_date=date, true_product=product))


def expand_dups(recs, dup_ids):
    out = []
    for r in recs:
        out.append(r)
        if r["id"] in dup_ids:
            d = copy.deepcopy(r)
            d["id"] = f"dup{r['id']}"
            out.append(d)
    return out


def inject_long(recs, file, sheet, C, lang, texts, plan):
    """Inject errors into long-format records; every record gets at most one error."""
    avail = list(range(len(recs)))
    rnd.shuffle(avail)
    dups = {}
    for kind, n in plan:
        for _ in range(n):
            r = recs[avail.pop()]
            v, date, prod = r["vals"], r["date"].isoformat(), PRODUCTS[r["p"]]["ru"]
            L = lambda col, et, o, i: log(file, sheet, r["id"], col, et, o, i, date, prod)
            if kind == "empty_qty":
                L(C["qty"], "empty_cell", v[C["qty"]], None); v[C["qty"]] = None
            elif kind == "empty_price":
                L(C["price"], "empty_cell", v[C["price"]], None); v[C["price"]] = None
            elif kind == "typo":
                new = make_typo(v[C["prod"]], lang)
                L(C["prod"], "typo_product_name", v[C["prod"]], new); v[C["prod"]] = new
            elif kind == "negative":
                L(C["qty"], "negative_quantity", v[C["qty"]], -v[C["qty"]]); v[C["qty"]] = -v[C["qty"]]
            elif kind == "invalid_date":
                L(C["date"], "invalid_date", v[C["date"]], C["bad_date"]); v[C["date"]] = C["bad_date"]
            elif kind == "date_format":
                new = r["date"].strftime("%Y/%m/%d")
                L(C["date"], "inconsistent_date_format", v[C["date"]], new); v[C["date"]] = new
            elif kind == "text_qty":
                new = rnd.choice(texts)
                L(C["qty"], "text_instead_of_number", v[C["qty"]], new); v[C["qty"]] = new
            elif kind == "price_comma":
                new = v[C["price"]].replace(".", ",")
                L(C["price"], "decimal_comma_in_price", v[C["price"]], new); v[C["price"]] = new
            elif kind == "duplicate":
                dups[r["id"]] = (date, prod)
    for rid, (date, prod) in dups.items():
        log(file, sheet, f"dup{rid}", "(entire row)", "duplicate_row", "-", "exact copy of previous row", date, prod)
    return expand_dups(recs, set(dups))


TRUTH = []

# ================================================================= 1) December 2022 - Azerbaijani
def build_dec():
    file, y, m = "sales_2022_12.xlsx", 2022, 12
    days, prices, qty = generate(y, m)
    TRUTH.append((m, truth_rows(y, m, days, prices, qty)))
    cols = ["Tarix", "Məhsul", "Kateqoriya", "Miqdar", "Qiymət"]
    recs = []
    for d in range(days):
        date = dt.date(y, m, d + 1)
        for i, p in enumerate(PRODUCTS):
            if qty[d][i] > 0:
                recs.append(dict(id=len(recs), date=date, p=i, vals={
                    "Tarix": date, "Məhsul": p["az"], "Kateqoriya": CATS[p["cat"]][1],
                    "Miqdar": int(qty[d][i]), "Qiymət": prices[i]}))
    C = dict(date="Tarix", prod="Məhsul", qty="Miqdar", price="Qiymət", bad_date="31.11.2022")
    plan = [("empty_qty", 1), ("empty_price", 1), ("typo", 2), ("negative", 1), ("duplicate", 1),
            ("invalid_date", 1), ("text_qty", 1), ("date_format", 1)]
    final = inject_long(recs, file, "Satış", C, "az", ["beş", "yoxdur", "N/A"], plan)
    ROWMAP[(file, "Satış")] = {r["id"]: i + 2 for i, r in enumerate(final)}

    wb = Workbook(); ws = wb.active; ws.title = "Satış"
    header(ws, 1, cols)
    for i, r in enumerate(final, 2):
        for c, name in enumerate(cols, 1):
            v = r["vals"][name]
            fmt = "DD.MM.YYYY" if isinstance(v, dt.date) else ("0.00" if name == "Qiymət" else None)
            put(ws, i, c, v, fmt=fmt)
    widths(ws, [13, 20, 24, 10, 10]); ws.freeze_panes = "A2"
    wb.save(OUT / file)


# ================================================================= 2) November 2022 - Russian
def build_nov():
    file, y, m = "sales_2022_11.xlsx", 2022, 11
    days, prices, qty = generate(y, m)
    TRUTH.append((m, truth_rows(y, m, days, prices, qty)))
    cols = ["Дата", "Наименование", "Кол-во", "Цена за ед.", "Сумма"]
    recs = []
    for d in range(days):
        date = dt.date(y, m, d + 1)
        order = list(range(N)); rnd.shuffle(order)  # this person lists items in receipt order
        for i in order:
            if qty[d][i] > 0:
                q = int(qty[d][i])
                recs.append(dict(id=len(recs), date=date, p=i, vals={
                    "Дата": date.isoformat(), "Наименование": PRODUCTS[i]["ru"], "Кол-во": q,
                    "Цена за ед.": f"{prices[i]:.2f} AZN", "Сумма": round(q * prices[i], 2)}))
    C = dict(date="Дата", prod="Наименование", qty="Кол-во", price="Цена за ед.", bad_date="2022-11-31")
    plan = [("empty_qty", 1), ("typo", 2), ("negative", 1), ("duplicate", 1), ("invalid_date", 1),
            ("text_qty", 1), ("price_comma", 1)]
    sheet = "Ноябрь"
    final = inject_long(recs, file, sheet, C, "ru", ["три", "нет", "н/д"], plan)
    ROWMAP[(file, sheet)] = {r["id"]: i + 5 for i, r in enumerate(final)}

    wb = Workbook(); ws = wb.active; ws.title = sheet
    titles = ["Мини-маркет «Хазар»", "Отчёт о продажах за ноябрь 2022 г.", "г. Баку, ул. Низами, 45  |  составил: администратор"]
    for r, t in enumerate(titles, 1):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
        put(ws, r, 1, t, bold=(r < 3), align="center")
    header(ws, 4, cols)
    for i, r in enumerate(final, 5):
        for c, name in enumerate(cols, 1):
            put(ws, i, c, r["vals"][name], fmt="0.00" if name == "Сумма" else None)
    widths(ws, [13, 24, 10, 14, 10])
    wb.save(OUT / file)


# ================================================================= 3) February 2022 - wide
def build_feb():
    file, y, m = "sales_2022_02.xlsx", 2022, 2
    days, prices, qty = generate(y, m)
    TRUTH.append((m, truth_rows(y, m, days, prices, qty)))
    rows = [dict(id=i, p=i, vals=[PRODUCTS[i]["az2"]] + [int(qty[d][i]) if d < days else None for d in range(31)])
            for i in range(N)]  # vals[0] = name, vals[d] = quantity on day d (1..31)
    sheet = "Satış"
    dup_row = rnd.randrange(N)
    typo_row = rnd.choice([i for i in range(N) if i != dup_row])
    cells = [(i, d) for i in range(N) if i != dup_row for d in range(1, days + 1) if rows[i]["vals"][d] > 0]
    rnd.shuffle(cells)

    def L(i, d, col, et, o, i_, date=None):
        log(file, sheet, rows[i]["id"], col, et, o, i_,
            date or dt.date(y, m, d).isoformat(), PRODUCTS[rows[i]["p"]]["ru"])

    for _ in range(2):  # empty cells
        i, d = cells.pop(); v = rows[i]["vals"]
        L(i, d, f"{get_column_letter(d + 1)} (gün {d})", "empty_cell", v[d], None); v[d] = None
    i, d = cells.pop(); v = rows[i]["vals"]  # negative
    L(i, d, f"{get_column_letter(d + 1)} (gün {d})", "negative_quantity", v[d], -v[d]); v[d] = -v[d]
    i, d = cells.pop(); v = rows[i]["vals"]  # text
    L(i, d, f"{get_column_letter(d + 1)} (gün {d})", "text_instead_of_number", v[d], "yoxdur"); v[d] = "yoxdur"
    i = rnd.choice([k for k in range(N) if k != dup_row]); d = rnd.choice([29, 30, 31])  # day that doesn't exist
    val = rnd.randint(3, 15)
    L(i, d, f"{get_column_letter(d + 1)} (gün {d})", "value_on_nonexistent_day", None, val,
      date="(day does not exist in February)"); rows[i]["vals"][d] = val
    new = make_typo(rows[typo_row]["vals"][0], "az")  # typo in row label
    L(typo_row, 0, "A (Məhsul)", "typo_product_name", rows[typo_row]["vals"][0], new,
      date="(whole row)"); rows[typo_row]["vals"][0] = new
    log(file, sheet, f"dup{dup_row}", "(entire row)", "duplicate_row", "-", "exact copy of previous row",
        "(whole row)", PRODUCTS[dup_row]["ru"])
    final = expand_dups(rows, {dup_row})
    ROWMAP[(file, sheet)] = {r["id"]: k + 2 for k, r in enumerate(final)}

    prices_rows = [dict(id=i, name=PRODUCTS[i]["az2"], price=prices[i]) for i in range(N)]
    j = rnd.randrange(N)
    log(file, "Qiymətlər", j, "B (Qiymət)", "empty_cell", prices[j], None, "(price list)", PRODUCTS[j]["ru"])
    prices_rows[j]["price"] = None
    ROWMAP[(file, "Qiymətlər")] = {i: i + 2 for i in range(N)}

    wb = Workbook(); ws = wb.active; ws.title = sheet
    put(ws, 1, 1, "Məhsul", True, align="center").fill = PatternFill("solid", fgColor="D9E1F2")
    for d in range(1, 32):
        put(ws, 1, d + 1, d, True, align="center").fill = PatternFill("solid", fgColor="D9E1F2")
    for k, r in enumerate(final, 2):
        for c, v in enumerate(r["vals"], 1):
            put(ws, k, c, v)
    widths(ws, [18] + [5] * 31); ws.freeze_panes = "B2"
    ws2 = wb.create_sheet("Qiymətlər")
    header(ws2, 1, ["Məhsul", "Qiymət"])
    for k, pr in enumerate(prices_rows, 2):
        put(ws2, k, 1, pr["name"]); put(ws2, k, 2, pr["price"], fmt="0.00")
    widths(ws2, [18, 10])
    wb.save(OUT / file)


# ================================================================= run
if __name__ == "__main__":
    build_feb(); build_nov(); build_dec()  # chronological order of truth: Feb, Nov, Dec

    gt = pd.DataFrame([row for _, rows in TRUTH for row in rows],
                      columns=["date", "product_name", "category", "quantity", "unit_price"])
    gt.to_csv(OUT / "ground_truth.csv", index=False, encoding="utf-8-sig")

    err = pd.DataFrame(ERRORS)
    err["excel_row"] = [ROWMAP[(e["file"], e["sheet"])][e["_id"]] for e in ERRORS]
    err = err[["file", "sheet", "excel_row", "column", "error_type", "original_value", "injected_value",
               "true_date", "true_product"]].sort_values(["file", "sheet", "excel_row"])
    err.to_csv(OUT / "errors_injected.csv", index=False, encoding="utf-8-sig")

    print(f"ground_truth.csv: {len(gt)} rows")
    print(err.groupby("file").size().to_string())
